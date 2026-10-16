from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import create_engine, func, select
from sqlalchemy.orm import Session

from backend.database import Base
from backend.main import (
    close_leader_group_buy,
    create_order,
    export_leader_group_buy_excel,
    get_leader_dashboard,
    get_leader_order,
    list_group_buy_products,
    list_group_buys,
    list_leader_orders,
)
from backend.models import GroupBuy, GroupBuyProduct, Order, OrderItem, Product, User
from backend.schemas import OrderCreate


class CloseAndExcelTest(unittest.TestCase):
    """全部使用隔离临时库；永不访问或截停正式团购 1。"""

    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        database_path = Path(self.temp_dir.name) / "close_export_test.db"
        self.engine = create_engine(f"sqlite:///{database_path.as_posix()}")
        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)
        self._seed()

    def tearDown(self) -> None:
        self.db.close()
        self.engine.dispose()
        self.temp_dir.cleanup()

    def _seed(self) -> None:
        self.current_user = User(id=1, openid="close-test-user", role="leader")
        self.db.add_all([
            self.current_user,
            GroupBuy(id=101, name="截团测试/秋日团", status="active", created_at=datetime(2026, 1, 1, 8), leader_id=1),
            GroupBuy(id=102, name="空白截团", status="closed", created_at=datetime(2026, 1, 1, 8), leader_id=1),
            GroupBuy(id=103, name="未知状态团", status="paused", created_at=datetime(2026, 1, 1, 8), leader_id=1),
            Product(id=201, name="商品A", description="A", emoji="A"),
            Product(id=202, name="商品B", description="B", emoji="B"),
        ])
        self.db.flush()
        self.db.add_all([
            GroupBuyProduct(group_buy_id=101, product_id=201, sale_price=Decimal("2.00"), display_order=1, is_active=True),
            GroupBuyProduct(group_buy_id=101, product_id=202, sale_price=Decimal("3.00"), display_order=2, is_active=True),
        ])
        moments = [datetime(2026, 1, day, 8) for day in (2, 3, 4, 4)]
        orders = [
            Order(id=301, order_no="TEST1", group_buy_id=101, group_buy_name="截团测试/秋日团", follow_no="001", customer_name="未填写", customer_phone=None, pickup_type=None, pickup_location=None, total_amount=Decimal("6.00"), status="submitted", created_at=moments[0], items=[OrderItem(product_name="商品A", unit_price=Decimal("2.00"), quantity=3, subtotal=Decimal("6.00"))]),
            Order(id=302, order_no="TEST2", group_buy_id=101, group_buy_name="截团测试/秋日团", follow_no="002", customer_name="甲", customer_phone="13000000000", pickup_type="dormitory", pickup_location="10", total_amount=Decimal("8.00"), status="submitted", created_at=moments[1], items=[OrderItem(product_name="商品A", unit_price=Decimal("2.50"), quantity=2, subtotal=Decimal("5.00")), OrderItem(product_name="商品B", unit_price=Decimal("3.00"), quantity=1, subtotal=Decimal("3.00"))]),
            Order(id=303, order_no="TEST3", group_buy_id=101, group_buy_name="截团测试/秋日团", follow_no="003", customer_name="乙", customer_phone="13100000000", pickup_type="dormitory", pickup_location="2", total_amount=Decimal("2.00"), status="submitted", created_at=moments[2], items=[OrderItem(product_name="商品A", unit_price=Decimal("2.00"), quantity=1, subtotal=Decimal("2.00"))]),
            Order(id=304, order_no="TEST4", group_buy_id=101, group_buy_name="截团测试/秋日团", follow_no="004", customer_name="丙", customer_phone="13200000000", pickup_type="other", pickup_location="图书馆", total_amount=Decimal("12.00"), status="submitted", created_at=moments[3], items=[OrderItem(product_name="商品B", unit_price=Decimal("3.00"), quantity=4, subtotal=Decimal("12.00"))]),
        ]
        self.db.add_all(orders)
        self.db.commit()

    def _order_payload(self) -> OrderCreate:
        return OrderCreate.model_validate({
            "group_buy_id": 101,
            "customer_name": "新同学",
            "customer_phone": "13800000000",
            "pickup_type": "dormitory",
            "pickup_location": "3",
            "items": [{"product_id": 201, "quantity": 2}],
        })

    def test_close_transitions_and_errors(self) -> None:
        result = close_leader_group_buy(101, self.db, self.current_user)
        self.assertEqual(result.model_dump(), {"status": "success", "group_buy_id": 101, "group_buy_status": "closed"})
        self.assertEqual(self.db.get(GroupBuy, 101).status, "closed")
        for group_id, detail in [(101, "团购已截团"), (103, "团购活动当前状态不可截团")]:
            with self.assertRaises(HTTPException) as error:
                close_leader_group_buy(group_id, self.db, self.current_user)
            self.assertEqual(error.exception.status_code, 409)
            self.assertEqual(error.exception.detail, detail)
        with self.assertRaises(HTTPException) as error:
            close_leader_group_buy(999, self.db, self.current_user)
        self.assertEqual(error.exception.status_code, 404)

    def test_closed_order_is_rejected_without_rows_or_follow_number_consumption(self) -> None:
        close_leader_group_buy(101, self.db, self.current_user)
        before_orders = self.db.scalar(select(func.count(Order.id)))
        before_items = self.db.scalar(select(func.count(OrderItem.id)))
        before_max = self.db.scalar(select(func.max(Order.follow_no)).where(Order.group_buy_id == 101))
        self.db.rollback()
        with self.assertRaises(HTTPException) as error:
            create_order(self._order_payload(), self.db, self.current_user)
        self.assertEqual(error.exception.status_code, 409)
        self.assertEqual(error.exception.detail, "团购活动当前不可下单")
        self.assertEqual(self.db.scalar(select(func.count(Order.id))), before_orders)
        self.assertEqual(self.db.scalar(select(func.count(OrderItem.id))), before_items)
        self.assertEqual(self.db.scalar(select(func.max(Order.follow_no)).where(Order.group_buy_id == 101)), before_max)

    def test_closed_group_leader_reads_remain_available(self) -> None:
        close_leader_group_buy(101, self.db, self.current_user)
        self.assertEqual(get_leader_dashboard(101, self.db, self.current_user).group_buy.status, "closed")
        self.assertEqual(len(list_leader_orders(101, self.db, self.current_user)), 4)
        self.assertEqual(get_leader_order(302, self.db, self.current_user).follow_no, "002")

    def test_export_guards(self) -> None:
        with self.assertRaises(HTTPException) as active_error:
            export_leader_group_buy_excel(101, self.db, self.current_user)
        self.assertEqual(active_error.exception.status_code, 409)
        self.assertEqual(active_error.exception.detail, "请先截团后再导出最终备货表")
        with self.assertRaises(HTTPException) as missing_error:
            export_leader_group_buy_excel(999, self.db, self.current_user)
        self.assertEqual(missing_error.exception.status_code, 404)

    def test_closed_export_is_valid_and_matches_all_aggregates(self) -> None:
        close_leader_group_buy(101, self.db, self.current_user)
        response = export_leader_group_buy_excel(101, self.db, self.current_user)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.media_type, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn(".xlsx", response.headers["content-disposition"])
        self.assertNotIn("/秋日团", response.headers["content-disposition"])

        workbook = load_workbook(BytesIO(response.body), data_only=True)
        self.assertEqual(workbook.sheetnames, ["商品汇总", "订单明细", "配送汇总"])
        products = workbook["商品汇总"]
        product_rows = list(products.iter_rows(min_row=7, max_row=8, max_col=3, values_only=True))
        self.assertEqual(product_rows, [("商品A", 6, 13), ("商品B", 5, 15)])
        self.assertEqual(tuple(products[9][index].value for index in range(3)), ("合计", 11, 28))

        details = workbook["订单明细"]
        detail_rows = list(details.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(detail_rows), 5)
        self.assertEqual(sum(1 for row in detail_rows if row[0] == "002"), 2)
        historical = next(row for row in detail_rows if row[0] == "001")
        self.assertEqual(historical[1], "未填写")
        self.assertEqual(historical[2], "未记录")
        self.assertEqual(historical[3], "未记录")
        dormitory = next(row for row in detail_rows if row[0] == "003")
        self.assertEqual(dormitory[3], "学生公寓2栋")

        pickups = list(workbook["配送汇总"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(pickups, [
            ("学生公寓", "2栋", 1, 1),
            ("学生公寓", "10栋", 1, 3),
            ("其他", "图书馆", 1, 4),
            ("历史未记录", "未记录", 1, 3),
        ])

    def test_empty_closed_group_exports_zero_totals(self) -> None:
        response = export_leader_group_buy_excel(102, self.db, self.current_user)
        workbook = load_workbook(BytesIO(response.body), data_only=True)
        products = workbook["商品汇总"]
        self.assertEqual(products["B4"].value, 0)
        self.assertEqual(products["D4"].value, 0)
        self.assertEqual(products["F4"].value, 0)
        self.assertEqual(tuple(products[7][index].value for index in range(3)), ("合计", 0, 0))
        self.assertEqual(workbook["订单明细"].max_row, 1)
        self.assertEqual(workbook["配送汇总"].max_row, 1)

    def test_active_regression_listing_products_pricing_and_follow_number(self) -> None:
        self.assertEqual([row.id for row in list_group_buys(False, self.db)], [101])
        self.assertEqual([row.id for row in list_group_buys(True, self.db)], [101, 102, 103])
        products = list_group_buy_products(101, self.db)
        self.assertEqual([row.product_id for row in products], [201, 202])
        created = create_order(self._order_payload(), self.db, self.current_user)
        self.assertEqual(created.follow_no, "005")
        self.assertEqual(created.total_amount, Decimal("4.00"))
        self.assertEqual(created.customer_name, "新同学")
        self.assertEqual(created.pickup_location, "3")

    def test_order_contact_and_pickup_validation_regression(self) -> None:
        invalid_payloads = [
            {"group_buy_id": 101, "customer_name": "", "customer_phone": "13800000000", "pickup_type": "dormitory", "pickup_location": "3", "items": [{"product_id": 201, "quantity": 1}]},
            {"group_buy_id": 101, "customer_name": "同学", "customer_phone": "微信号abc", "pickup_type": "dormitory", "pickup_location": "3", "items": [{"product_id": 201, "quantity": 1}]},
            {"group_buy_id": 101, "customer_name": "同学", "customer_phone": "13800000000", "pickup_type": "dormitory", "pickup_location": "三", "items": [{"product_id": 201, "quantity": 1}]},
            {"group_buy_id": 101, "customer_name": "同学", "customer_phone": "13800000000", "pickup_type": "other", "pickup_location": "北", "items": [{"product_id": 201, "quantity": 1}]},
        ]
        for payload in invalid_payloads:
            with self.assertRaises(ValidationError):
                OrderCreate.model_validate(payload)


if __name__ == "__main__":
    unittest.main()

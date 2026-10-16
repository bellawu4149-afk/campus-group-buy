from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import create_engine, event, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from backend.database import Base
from backend.main import (
    clone_leader_group_buy,
    create_order,
    export_leader_group_buy_excel,
    get_leader_dashboard,
)
from backend.models import GroupBuy, GroupBuyProduct, Order, OrderItem, Product, User
from backend.schemas import LeaderGroupBuyCloneCreate, OrderCreate


class UserRoleFoundationTest(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        database_path = Path(self.temp_dir.name) / "user_role_test.db"
        self.engine = create_engine(f"sqlite:///{database_path.as_posix()}")

        @event.listens_for(self.engine, "connect")
        def enable_foreign_keys(dbapi_connection, connection_record):
            dbapi_connection.execute("PRAGMA foreign_keys=ON")

        Base.metadata.create_all(self.engine)
        self.db = Session(self.engine)
        self._seed_business_data()

    def tearDown(self) -> None:
        self.db.close()
        self.engine.dispose()
        self.temp_dir.cleanup()

    def _seed_business_data(self) -> None:
        group_buy = GroupBuy(
            id=501,
            name="历史无归属团",
            status="active",
            created_at=datetime(2026, 1, 1, 8),
        )
        product = Product(id=601, name="测试商品", description="", emoji="T")
        self.db.add_all([group_buy, product])
        self.db.flush()
        self.db.add(
            GroupBuyProduct(
                group_buy_id=501,
                product_id=601,
                sale_price=Decimal("9.90"),
                display_order=1,
                is_active=True,
                last_week_sales=0,
            )
        )
        self.db.add(
            Order(
                id=701,
                order_no="HISTORY001",
                group_buy_id=501,
                group_buy_name="历史无归属团",
                follow_no="001",
                customer_name="历史用户",
                customer_phone=None,
                pickup_type=None,
                pickup_location=None,
                total_amount=Decimal("9.90"),
                status="submitted",
                created_at=datetime(2026, 1, 2, 8),
                items=[OrderItem(product_name="测试商品", unit_price=Decimal("9.90"), quantity=1, subtotal=Decimal("9.90"))],
            )
        )
        self.db.commit()

    def _order_payload(self, **extra) -> OrderCreate:
        payload = {
            "group_buy_id": 501,
            "customer_name": "无登录学生",
            "customer_phone": "13800000000",
            "pickup_type": "dormitory",
            "pickup_location": "3",
            "items": [{"product_id": 601, "quantity": 1}],
        }
        payload.update(extra)
        return OrderCreate.model_validate(payload)

    def test_users_roles_defaults_unique_and_nullable_openid(self) -> None:
        student = User(openid="openid-student", nickname="学生")
        leader = User(openid="openid-leader", nickname="团长", role="leader")
        anonymous_one = User(openid=None, nickname=None)
        anonymous_two = User(openid=None, nickname="待绑定")
        self.db.add_all([student, leader, anonymous_one, anonymous_two])
        self.db.commit()
        self.assertEqual(student.role, "student")
        self.assertEqual(leader.role, "leader")
        self.assertIsNotNone(student.created_at)
        self.assertEqual(self.db.scalar(select(func.count(User.id))), 4)

        self.db.add(User(openid="openid-student", nickname="重复"))
        with self.assertRaises(IntegrityError):
            self.db.commit()
        self.db.rollback()

    def test_user_order_and_leader_group_relationships(self) -> None:
        student = User(openid="relationship-student", role="student")
        leader = User(openid="relationship-leader", role="leader")
        order = self.db.get(Order, 701)
        group_buy = self.db.get(GroupBuy, 501)
        order.user = student
        group_buy.leader = leader
        self.db.commit()
        self.assertEqual(student.orders[0].id, 701)
        self.assertEqual(order.user.id, student.id)
        self.assertEqual(leader.group_buys[0].id, 501)
        self.assertEqual(group_buy.leader.id, leader.id)

    def test_null_ownership_history_is_readable_and_exportable(self) -> None:
        order = self.db.get(Order, 701)
        group_buy = self.db.get(GroupBuy, 501)
        self.assertIsNone(order.user_id)
        self.assertIsNone(order.user)
        self.assertIsNone(group_buy.leader_id)
        self.assertIsNone(group_buy.leader)
        leader = User(openid="history-leader", role="leader")
        self.db.add(leader)
        self.db.commit()
        with self.assertRaises(HTTPException):
            get_leader_dashboard(501, self.db, leader)
        group_buy.leader_id = leader.id
        group_buy.status = "closed"
        self.db.commit()
        self.assertEqual(get_leader_dashboard(501, self.db, leader).summary.order_count, 1)
        response = export_leader_group_buy_excel(501, self.db, leader)
        workbook = load_workbook(BytesIO(response.body), data_only=True)
        self.assertEqual(workbook["订单明细"].max_row, 2)

    def test_deleting_owned_user_does_not_cascade_business_rows(self) -> None:
        user = User(openid="protected-user", role="leader")
        self.db.add(user)
        self.db.flush()
        self.db.get(Order, 701).user_id = user.id
        self.db.get(GroupBuy, 501).leader_id = user.id
        self.db.commit()
        self.db.delete(user)
        with self.assertRaises(IntegrityError):
            self.db.commit()
        self.db.rollback()
        self.assertIsNotNone(self.db.get(User, user.id))
        self.assertIsNotNone(self.db.get(Order, 701))
        self.assertIsNotNone(self.db.get(GroupBuy, 501))

    def test_clients_cannot_forge_ownership_fields(self) -> None:
        with self.assertRaises(ValidationError):
            self._order_payload(user_id=123)
        with self.assertRaises(ValidationError):
            LeaderGroupBuyCloneCreate.model_validate({"name": "新团", "leader_id": 123})

    def test_authenticated_order_and_clone_bind_current_user(self) -> None:
        current_user = User(openid="ordering-user", role="leader")
        self.db.add(current_user)
        self.db.commit()
        created_order = create_order(self._order_payload(), self.db, current_user)
        order = self.db.get(Order, created_order.order_id)
        self.assertEqual(order.user_id, current_user.id)
        self.db.get(GroupBuy, 501).leader_id = current_user.id
        self.db.commit()
        cloned = clone_leader_group_buy(
            501,
            LeaderGroupBuyCloneCreate(name="未登录复制团"),
            self.db,
            current_user,
        )
        cloned_group = self.db.get(GroupBuy, cloned.group_buy.id)
        self.assertEqual(cloned_group.leader_id, current_user.id)


if __name__ == "__main__":
    unittest.main()
